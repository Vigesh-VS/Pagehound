import io
import shutil

import openpyxl
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from app.extractors import UnsupportedType, extract
from app.pdf_utils import ExtractError


def csv_bytes(rows=45, delimiter=","):
    lines = [delimiter.join(["product", "price", "city"])]
    lines += [delimiter.join([f"item{i}", str(i * 10), "Chennai" if i % 2 else "Pune"]) for i in range(1, rows + 1)]
    return "\n".join(lines).encode()


def test_csv_gets_a_summary_and_row_groups():
    result = extract("sales.csv", csv_bytes())
    labels = [u.label for u in result.units]
    assert result.kind == "table"
    assert labels[0] == "Summary" and labels[1] == "Rows 1-20" and labels[-1] == "Rows 41-45"
    summary = result.units[0].text
    assert "45 data rows" in summary and "price (numeric)" in summary and "sum 10350" in summary
    assert "Chennai (23)" in summary


def test_tsv_and_semicolon_csv_are_detected():
    assert extract("a.tsv", csv_bytes(10, "\t")).units[1].text.startswith("Row 1: product: item1")
    assert extract("a.csv", csv_bytes(10, ";")).units[1].text.startswith("Row 1: product: item1")


def test_table_without_data_is_rejected():
    with pytest.raises(ExtractError):
        extract("empty.csv", b"only,a,header\n")


def test_only_the_first_rows_are_searchable_but_stats_cover_all(monkeypatch):
    from app import config

    monkeypatch.setattr(config, "EMBED_ROWS", 30)
    result = extract("big.csv", csv_bytes(100))
    summary = result.units[0].text
    assert result.truncated is True
    assert "100 data rows" in summary and "first 30 rows are searchable" in summary
    assert result.units[-1].label == "Rows 21-30"


def test_read_cap_limits_the_rows_used_for_statistics(monkeypatch):
    from app import config

    monkeypatch.setattr(config, "MAX_ROWS", 50)
    result = extract("huge.csv", csv_bytes(100))
    assert result.truncated is True and "50 data rows" in result.units[0].text


def test_xlsx_labels_include_the_sheet_name():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["item", "cost"])
    for i in range(1, 6):
        sheet.append([f"thing{i}", i * 100])
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = extract("budget.xlsx", buffer.getvalue())
    assert [u.label for u in result.units] == ["Budget summary", "Budget rows 1-5"]
    assert "cost (numeric)" in result.units[0].text


def test_docx_text_and_tables_are_read_in_order():
    document = Document()
    document.add_paragraph("Quarterly report introduction.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Sales"
    table.cell(1, 0).text, table.cell(1, 1).text = "North", "500"
    document.add_paragraph("Closing remarks.")
    buffer = io.BytesIO()
    document.save(buffer)
    result = extract("report.docx", buffer.getvalue())
    text = result.units[0].text
    assert result.kind == "document" and result.units[0].label == "Part 1"
    assert text.index("introduction") < text.index("North | 500") < text.index("Closing")


def test_pptx_slides_and_notes():
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Roadmap"
    first.placeholders[1].text = "Launch in March"
    first.notes_slide.notes_text_frame.text = "Mention the pilot"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    buffer = io.BytesIO()
    presentation.save(buffer)
    result = extract("deck.pptx", buffer.getvalue())
    assert [u.label for u in result.units] == ["Slide 1", "Slide 2"]
    assert "Launch in March" in result.units[0].text and "Mention the pilot" in result.units[0].text


def test_plain_text_is_split_into_parts(monkeypatch):
    from app import config

    monkeypatch.setattr(config, "TEXT_UNIT_CHARS", 100)
    body = "\n".join(f"Paragraph number {i} with some filler words." for i in range(12))
    result = extract("notes.txt", body.encode())
    assert len(result.units) > 1 and result.units[0].label == "Part 1"


def test_latin1_text_does_not_crash():
    assert "café" in extract("old.txt", "café au lait".encode("latin-1")).units[0].text


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="Tesseract is not installed")
def test_image_is_read_with_ocr():
    image = Image.new("RGB", (900, 260), "white")
    ImageDraw.Draw(image).text((30, 90), "Invoice total 4521 rupees", fill="black", font=ImageFont.load_default(size=44))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    result = extract("photo.png", buffer.getvalue())
    assert result.kind == "image" and result.units[0].ocr is True and result.units[0].label == "Image"
    assert "invoice" in result.units[0].text.lower()


def test_broken_image_and_unsupported_type():
    with pytest.raises(ExtractError):
        extract("photo.png", b"not an image")
    with pytest.raises(UnsupportedType):
        extract("movie.mp4", b"data")
    with pytest.raises(UnsupportedType):
        extract("old.doc", b"data")


def test_corrupt_office_file_is_rejected():
    with pytest.raises(ExtractError):
        extract("bad.docx", b"this is not a zip")
