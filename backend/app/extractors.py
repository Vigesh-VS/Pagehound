import csv
import io
import zipfile
from collections import Counter
from dataclasses import dataclass

import openpyxl
from docx import Document
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps
from pptx import Presentation

from . import config
from .pdf_utils import ExtractError, PageText, extract_pages, ocr_image

IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED = "PDF, images (PNG, JPG, WEBP, BMP, TIFF), CSV, TSV, XLSX, DOCX, PPTX, TXT and MD"
MAX_UNPACKED_BYTES = 150 * 1024 * 1024


class UnsupportedType(ExtractError):
    pass


@dataclass
class Extracted:
    units: list
    total_units: int
    kind: str
    truncated: bool = False


def file_extension(filename):
    name = filename.lower()
    return name[name.rfind("."):] if "." in name else ""


def _check_zip(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise ExtractError("This file is too large once unpacked.")
    except zipfile.BadZipFile as exc:
        raise ExtractError("This file looks corrupted or is not a valid Office file.") from exc


def _decode(data):
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _fmt(number):
    return f"{number:.2f}".rstrip("0").rstrip(".")


def _number(value):
    cleaned = value.replace(",", "").replace("$", "").replace("₹", "").replace("%", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def _summary(header, rows):
    lines = [
        f"Table with {len(rows)} data rows and {len(header)} columns.",
        "Columns: " + ", ".join(header) + ".",
    ]
    for index, name in enumerate(header):
        values = [row[index] for row in rows if index < len(row) and row[index].strip()]
        if not values:
            continue
        numbers = [n for n in (_number(v) for v in values) if n is not None]
        if len(numbers) >= 0.8 * len(values):
            lines.append(
                f"{name} (numeric): count {len(numbers)}, min {_fmt(min(numbers))}, max {_fmt(max(numbers))}, "
                f"mean {_fmt(sum(numbers) / len(numbers))}, sum {_fmt(sum(numbers))}."
            )
        else:
            common = Counter(values).most_common(5)
            listed = ", ".join(f"{value} ({count})" for value, count in common)
            lines.append(f"{name} (text): {len(set(values))} distinct values, most common: {listed}.")
    return "\n".join(lines)


def _table_units(sheet, header, rows):
    header = [h.strip() or f"column{i + 1}" for i, h in enumerate(header)]
    searchable = rows[: config.EMBED_ROWS]
    summary = _summary(header, rows)
    if len(rows) > len(searchable):
        summary += (
            f"\nOnly the first {len(searchable)} rows are searchable; "
            f"the statistics above cover all {len(rows)} rows."
        )
    units = [(f"{sheet} summary" if sheet else "Summary", summary)]
    step = config.ROWS_PER_UNIT
    for start in range(0, len(searchable), step):
        group = searchable[start:start + step]
        lines = []
        for offset, row in enumerate(group):
            cells = [f"{h}: {v}" for h, v in zip(header, row) if v.strip()]
            lines.append(f"Row {start + offset + 1}: " + " | ".join(cells))
        span = f"{start + 1}-{start + len(group)}"
        label = f"{sheet} rows {span}" if sheet else f"Rows {span}"
        units.append((label, "\n".join(lines)))
    return units


def _number_units(units, ocr=False):
    return [PageText(i + 1, text, ocr, label) for i, (label, text) in enumerate(units)]


def _text_units(text):
    paragraphs = [p.strip() for p in text.splitlines() if p.strip()]
    groups, current, length = [], [], 0
    for paragraph in paragraphs:
        if current and length + len(paragraph) > config.TEXT_UNIT_CHARS:
            groups.append("\n".join(current))
            current, length = [], 0
        current.append(paragraph)
        length += len(paragraph) + 1
    if current:
        groups.append("\n".join(current))
    return [(f"Part {i + 1}", group) for i, group in enumerate(groups)]


def _capped(units, kind):
    total = len(units)
    return Extracted(_number_units(units[: config.MAX_UNITS]), total, kind, total > config.MAX_UNITS)


def _pdf(data, max_pages=None):
    pages, total = extract_pages(data, max_pages)
    return Extracted(pages, total, "pdf", total > len(pages))


def _image(data):
    try:
        image = Image.open(io.BytesIO(data))
        if image.width * image.height > config.MAX_IMAGE_PIXELS:
            raise ExtractError("This image is too large. Please upload a smaller one.")
        image.load()
    except ExtractError:
        raise
    except Exception as exc:
        raise ExtractError("Could not read this file as an image.") from exc

    gray = ImageOps.autocontrast(ImageOps.grayscale(ImageOps.exif_transpose(image)))
    if gray.width < 1500:
        scale = min(3.0, 1500 / gray.width)
        gray = gray.resize((int(gray.width * scale), int(gray.height * scale)), Image.LANCZOS)
    return Extracted([PageText(1, ocr_image(gray), True, "Image")], 1, "image")


def _plain_text(data):
    return _capped(_text_units(_decode(data)), "text")


def _delimited(data, extension):
    text = _decode(data)
    delimiter = "\t" if extension == ".tsv" else ","
    if extension == ".csv":
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            pass
    rows, truncated = [], False
    try:
        for row in csv.reader(io.StringIO(text), delimiter=delimiter):
            if any(cell.strip() for cell in row):
                rows.append(row)
            if len(rows) > config.MAX_ROWS + 1:
                truncated = True
                break
    except csv.Error as exc:
        raise ExtractError("Could not parse this file as a table.") from exc
    if len(rows) < 2:
        raise ExtractError("This table needs a header row and at least one data row.")
    data_rows = rows[1 : config.MAX_ROWS + 1]
    result = _capped(_table_units("", rows[0], data_rows), "table")
    result.truncated = result.truncated or truncated or len(data_rows) > config.EMBED_ROWS
    return result


def _xlsx(data):
    _check_zip(data)
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractError("Could not read this Excel file.") from exc

    units, truncated = [], False
    for sheet in workbook.worksheets[:10]:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell) for cell in row]
            if any(cell.strip() for cell in cells):
                rows.append(cells)
            if len(rows) > config.MAX_ROWS + 1:
                truncated = True
                break
        if len(rows) >= 2:
            data_rows = rows[1 : config.MAX_ROWS + 1]
            truncated = truncated or len(data_rows) > config.EMBED_ROWS
            units += _table_units(sheet.title, rows[0], data_rows)
    workbook.close()

    if not units:
        raise ExtractError("No tables with data were found in this workbook.")
    result = _capped(units, "table")
    result.truncated = result.truncated or truncated
    return result


def _docx(data):
    _check_zip(data)
    try:
        document = Document(io.BytesIO(data))
        lines = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                lines.append(block.text)
            else:
                for row in block.rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    except Exception as exc:
        raise ExtractError("Could not read this Word file.") from exc
    return _capped(_text_units("\n".join(lines)), "document")


def _pptx(data):
    _check_zip(data)
    try:
        presentation = Presentation(io.BytesIO(data))
        units = []
        for number, slide in enumerate(presentation.slides, 1):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text.strip())
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
            units.append((f"Slide {number}", "\n".join(parts)))
    except Exception as exc:
        raise ExtractError("Could not read this PowerPoint file.") from exc
    return _capped(units, "slides")


def extract(filename, data, max_pages=None):
    extension = file_extension(filename)
    if extension == ".pdf" or (not extension and data.startswith(b"%PDF-")):
        return _pdf(data, max_pages)
    if extension in IMAGE_TYPES:
        return _image(data)
    if extension in {".csv", ".tsv"}:
        return _delimited(data, extension)
    if extension == ".xlsx":
        return _xlsx(data)
    if extension == ".docx":
        return _docx(data)
    if extension == ".pptx":
        return _pptx(data)
    if extension in {".txt", ".md"}:
        return _plain_text(data)
    raise UnsupportedType(f"This file type isn't supported. Supported types: {SUPPORTED}.")
